"""Testes de integração para o gerador."""

from pathlib import Path

from docx import Document
from openpyxl import Workbook

from gerador_peticoes.config import Config
from gerador_peticoes.gerador import executar


def _criar_modelo(path: Path, texto: str = "Eu, {Nome}, gênero {Genero}.") -> None:
    doc = Document()
    doc.add_paragraph(texto)
    doc.save(str(path))


def _criar_planilha(path: Path, registros: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    if registros:
        headers = list(registros[0].keys())
        ws.append(headers)
        for reg in registros:
            ws.append([reg[h] for h in headers])
    wb.save(str(path))


class TestExecutar:
    def test_fluxo_basico(self, tmp_path):
        entrada = tmp_path / "entrada"
        saida = tmp_path / "saida"
        entrada.mkdir()

        _criar_modelo(entrada / "modelo_masculino.docx")
        _criar_modelo(entrada / "modelo_feminino.docx")
        _criar_planilha(entrada / "dados_clientes.xlsx", [
            {"Nome": "João", "Genero": "M"},
            {"Nome": "Maria", "Genero": "F"},
        ])

        cfg = Config(dir_entrada=entrada, dir_saida=saida)
        result = executar(cfg)

        assert result == 0
        assert (saida / "Peticao - João.docx").exists()
        assert (saida / "Peticao - Maria.docx").exists()
        assert (saida / "relatorio_conferencia.xlsx").exists()
        assert (saida / "log_execucao.txt").exists()

    def test_dry_run_nao_cria_docx(self, tmp_path):
        entrada = tmp_path / "entrada"
        saida = tmp_path / "saida"
        entrada.mkdir()

        _criar_modelo(entrada / "modelo_masculino.docx")
        _criar_modelo(entrada / "modelo_feminino.docx")
        _criar_planilha(entrada / "dados_clientes.xlsx", [
            {"Nome": "João", "Genero": "M"},
        ])

        cfg = Config(dir_entrada=entrada, dir_saida=saida, dry_run=True)
        result = executar(cfg)

        assert result == 0
        assert not (saida / "Peticao - João.docx").exists()

    def test_genero_invalido_gera_aviso(self, tmp_path):
        entrada = tmp_path / "entrada"
        saida = tmp_path / "saida"
        entrada.mkdir()

        _criar_modelo(entrada / "modelo_masculino.docx")
        _criar_modelo(entrada / "modelo_feminino.docx")
        _criar_planilha(entrada / "dados_clientes.xlsx", [
            {"Nome": "Alex", "Genero": "X"},
        ])

        cfg = Config(dir_entrada=entrada, dir_saida=saida)
        result = executar(cfg)

        assert result == 0
        assert not (saida / "Peticao - Alex.docx").exists()

    def test_arquivo_entrada_faltando(self, tmp_path):
        cfg = Config(dir_entrada=tmp_path / "nao_existe", dir_saida=tmp_path / "saida")
        result = executar(cfg)
        assert result == 1

    def test_nome_execucao_no_log(self, tmp_path):
        entrada = tmp_path / "entrada"
        saida = tmp_path / "saida" / "saida_10-03-2026_14h30m00s"
        entrada.mkdir()

        _criar_modelo(entrada / "modelo_masculino.docx")
        _criar_modelo(entrada / "modelo_feminino.docx")
        _criar_planilha(entrada / "dados_clientes.xlsx", [
            {"Nome": "João", "Genero": "M"},
        ])

        cfg = Config(
            dir_entrada=entrada,
            dir_saida=saida,
            nome_execucao="saida_10-03-2026_14h30m00s",
        )
        result = executar(cfg)

        assert result == 0
        assert saida.exists()
        assert (saida / "Peticao - João.docx").exists()

        log_texto = (saida / "log_execucao.txt").read_text(encoding="utf-8")
        assert "saida_10-03-2026_14h30m00s" in log_texto

    def test_workers_paralelo(self, tmp_path):
        entrada = tmp_path / "entrada"
        saida = tmp_path / "saida"
        entrada.mkdir()

        _criar_modelo(entrada / "modelo_masculino.docx")
        _criar_modelo(entrada / "modelo_feminino.docx")
        _criar_planilha(entrada / "dados_clientes.xlsx", [
            {"Nome": "João", "Genero": "M"},
            {"Nome": "Maria", "Genero": "F"},
            {"Nome": "Pedro", "Genero": "M"},
        ])

        cfg = Config(dir_entrada=entrada, dir_saida=saida, workers=2)
        result = executar(cfg)

        assert result == 0
        assert (saida / "Peticao - João.docx").exists()
        assert (saida / "Peticao - Maria.docx").exists()
        assert (saida / "Peticao - Pedro.docx").exists()

    def test_variaveis_derivadas_substituidas(self, tmp_path):
        """Variáveis derivadas (ListaPreservacoes, ListaRTs, etc.) são
        substituídas corretamente no documento final."""
        entrada = tmp_path / "entrada"
        saida = tmp_path / "saida"
        entrada.mkdir()

        # Modelo que usa variáveis derivadas
        texto_modelo = (
            "Eu, {Nome}, preservei {ListaPreservacoes}. "
            "RT: {ListaRTs}. Data: {DataPeticao}."
        )
        _criar_modelo(entrada / "modelo_masculino.docx", texto_modelo)
        _criar_modelo(entrada / "modelo_feminino.docx", texto_modelo)
        _criar_planilha(entrada / "dados_clientes.xlsx", [
            {
                "Nome": "Almir",
                "Genero": "M",
                "PreservacaoSP": "07/2012,R$7.487,54",
                "RT_Indenizatoria": "0000001-61.2022.5.10.0017",
            },
        ])

        cfg = Config(dir_entrada=entrada, dir_saida=saida)
        result = executar(cfg)
        assert result == 0

        doc = Document(str(saida / "Peticao - Almir.docx"))
        texto = doc.paragraphs[0].text
        assert "julho de 2012" in texto
        assert "R$ 7.487,54" in texto
        assert "0000001-61.2022.5.10.0017" in texto
        assert "{ListaPreservacoes}" not in texto
        assert "{ListaRTs}" not in texto
        assert "{DataPeticao}" not in texto

    def test_bloco_condicional_removido_sem_preservacao(self, tmp_path):
        """Item 4.2 é removido quando cliente não tem PreservacaoSP."""
        entrada = tmp_path / "entrada"
        saida = tmp_path / "saida"
        entrada.mkdir()

        doc = Document()
        doc.add_paragraph("4.1. Primeiro pedido de {Nome}.")
        doc.add_paragraph("{SE:PreservacaoSP}")
        doc.add_paragraph("4.2. Preservou salário {ListaPreservacoes}.")
        doc.add_paragraph("{FIM_SE:PreservacaoSP}")
        doc.add_paragraph("5. Pedido final.")
        doc.save(str(entrada / "modelo_feminino.docx"))
        doc.save(str(entrada / "modelo_masculino.docx"))

        _criar_planilha(entrada / "dados_clientes.xlsx", [
            {"Nome": "Roseli", "Genero": "F", "PreservacaoSP": ""},
        ])

        cfg = Config(dir_entrada=entrada, dir_saida=saida)
        result = executar(cfg)
        assert result == 0

        doc_resultado = Document(str(saida / "Peticao - Roseli.docx"))
        textos = [p.text for p in doc_resultado.paragraphs]
        assert any("Roseli" in t for t in textos)
        assert not any("4.2" in t for t in textos)
        assert not any("ListaPreservacoes" in t for t in textos)
        assert any("5. Pedido final." in t for t in textos)

    def test_bloco_condicional_mantido_com_preservacao(self, tmp_path):
        """Item 4.2 é mantido e formatado quando cliente tem PreservacaoSP."""
        entrada = tmp_path / "entrada"
        saida = tmp_path / "saida"
        entrada.mkdir()

        doc = Document()
        doc.add_paragraph("4.1. Primeiro pedido de {Nome}.")
        doc.add_paragraph("{SE:PreservacaoSP}")
        doc.add_paragraph("4.2. Preservou salário {ListaPreservacoes}.")
        doc.add_paragraph("{FIM_SE:PreservacaoSP}")
        doc.add_paragraph("5. Pedido final.")
        doc.save(str(entrada / "modelo_feminino.docx"))
        doc.save(str(entrada / "modelo_masculino.docx"))

        _criar_planilha(entrada / "dados_clientes.xlsx", [
            {"Nome": "Almir", "Genero": "M", "PreservacaoSP": "07/2012,R$7.487,54"},
        ])

        cfg = Config(dir_entrada=entrada, dir_saida=saida)
        result = executar(cfg)
        assert result == 0

        doc_resultado = Document(str(saida / "Peticao - Almir.docx"))
        textos = [p.text for p in doc_resultado.paragraphs]
        texto_42 = [t for t in textos if "4.2" in t]
        assert len(texto_42) == 1
        assert "julho de 2012" in texto_42[0]
        assert "{ListaPreservacoes}" not in texto_42[0]
