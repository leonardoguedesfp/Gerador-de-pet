"""Testes para o módulo de documento."""

import pytest

from gerador_peticoes.documento import nome_arquivo_seguro, gerar_peticao

from docx import Document


class TestNomeArquivoSeguro:
    def test_remove_caracteres_proibidos(self):
        assert nome_arquivo_seguro('João <Silva>') == "João Silva"

    def test_remove_aspas_e_barras(self):
        assert nome_arquivo_seguro('Maria "Ana" S/A') == "Maria Ana SA"

    def test_strip(self):
        assert nome_arquivo_seguro("  José  ") == "José"

    def test_nome_limpo(self):
        assert nome_arquivo_seguro("Carlos Santos") == "Carlos Santos"


class TestGerarPeticao:
    def test_substituicao_basica(self, tmp_path):
        # Cria modelo
        modelo = tmp_path / "modelo.docx"
        doc = Document()
        doc.add_paragraph("Eu, {Nome}, portador do CPF {CPF}.")
        doc.save(str(modelo))

        saida = tmp_path / "saida" / "peticao.docx"
        gerar_peticao(modelo, {"Nome": "João Silva", "CPF": "123.456.789-00"}, saida)

        assert saida.exists()
        doc_resultado = Document(str(saida))
        texto = doc_resultado.paragraphs[0].text
        assert "João Silva" in texto
        assert "123.456.789-00" in texto
        assert "{Nome}" not in texto

    def test_variavel_nao_encontrada_e_removida(self, tmp_path):
        modelo = tmp_path / "modelo.docx"
        doc = Document()
        doc.add_paragraph("Valor: {ValorInexistente}")
        doc.save(str(modelo))

        saida = tmp_path / "peticao.docx"
        gerar_peticao(modelo, {"Nome": "Test"}, saida)

        doc_resultado = Document(str(saida))
        texto = doc_resultado.paragraphs[0].text
        assert "{ValorInexistente}" not in texto
        assert "Valor:" in texto
