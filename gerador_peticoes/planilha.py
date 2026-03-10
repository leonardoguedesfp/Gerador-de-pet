"""Leitura de planilha Excel com formatação inteligente."""

from pathlib import Path

from openpyxl import load_workbook

from .formatacao import formatar_valor_celula


def ler_planilha(
    caminho: Path,
    colunas_monetarias: list[str] | None = None,
    formato_data: str = "%d/%m/%Y",
) -> tuple[list[str], list[dict[str, str]]]:
    """
    Lê a planilha Excel e retorna (cabeçalhos, lista_de_registros).

    Cada registro é um dict {coluna: valor_formatado}.
    Linhas totalmente vazias são ignoradas.
    """
    wb = load_workbook(caminho, data_only=True)
    ws = wb.active

    cabecalhos = [
        str(cell.value).strip() if cell.value is not None else ""
        for cell in ws[1]
    ]

    dados: list[dict[str, str]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in row):
            continue
        registro = {}
        for i, valor in enumerate(row):
            if i < len(cabecalhos) and cabecalhos[i]:
                nome_col = cabecalhos[i]
                registro[nome_col] = formatar_valor_celula(
                    nome_col, valor, colunas_monetarias, formato_data,
                )
        dados.append(registro)

    wb.close()
    return cabecalhos, dados
