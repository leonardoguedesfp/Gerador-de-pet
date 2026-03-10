"""Geração de relatório de conferência em Excel."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .logger import Logger


def gerar_relatorio_conferencia(
    resultados: list[dict],
    caminho: Path,
    log: Logger,
    colunas_extras: list[str] | None = None,
) -> None:
    """
    Gera planilha Excel de conferência com o resultado de cada petição.

    Colunas: Nº, Nome, Gênero, [extras], Arquivo Gerado, Status, Observação.
    """
    colunas_extras = colunas_extras or []

    wb = Workbook()
    ws = wb.active
    ws.title = "Conferência"

    colunas_header = ["Nº", "Nome", "Gênero"]
    colunas_header.extend(colunas_extras)
    colunas_header.extend(["Arquivo Gerado", "Status", "Observação"])

    # Estilos
    fonte_header = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    alinhamento_centro = Alignment(horizontal="center", vertical="center")
    alinhamento_esquerda = Alignment(horizontal="left", vertical="center", wrap_text=True)
    borda = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    estilos_status = {
        "OK": (
            PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"),
            Font(name="Arial", size=10, color="2E7D32", bold=True),
        ),
        "ERRO": (
            PatternFill(start_color="FCE4EC", end_color="FCE4EC", fill_type="solid"),
            Font(name="Arial", size=10, color="C62828", bold=True),
        ),
        "AVISO": (
            PatternFill(start_color="FFF3E0", end_color="FFF3E0", fill_type="solid"),
            Font(name="Arial", size=10, color="E65100", bold=True),
        ),
    }

    # Cabeçalho
    for col_idx, titulo in enumerate(colunas_header, start=1):
        cell = ws.cell(row=1, column=col_idx, value=titulo)
        cell.font = fonte_header
        cell.fill = fill_header
        cell.alignment = alinhamento_centro
        cell.border = borda

    # Dados
    for row_idx, r in enumerate(resultados, start=2):
        valores = [r.get("numero", ""), r.get("nome", ""), r.get("genero", "")]
        for col_extra in colunas_extras:
            valores.append(r.get(col_extra, ""))
        valores.extend([r.get("arquivo", ""), r.get("status", ""), r.get("observacao", "")])

        status = r.get("status", "")
        for col_idx, valor in enumerate(valores, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=valor)
            cell.border = borda
            cell.font = Font(name="Arial", size=10)
            cell.alignment = alinhamento_esquerda

            col_nome = colunas_header[col_idx - 1]
            if col_nome == "Status" and status in estilos_status:
                fill, fonte = estilos_status[status]
                cell.fill = fill
                cell.font = fonte
                cell.alignment = alinhamento_centro
            elif col_nome in ("Nº", "Gênero"):
                cell.alignment = alinhamento_centro

    # Largura das colunas
    larguras = {"Nº": 6, "Nome": 35, "Gênero": 10, "Arquivo Gerado": 45,
                "Status": 10, "Observação": 40}
    for col_idx, titulo in enumerate(colunas_header, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = larguras.get(titulo, 20)

    # Resumo
    row_resumo = len(resultados) + 3
    total_ok = sum(1 for r in resultados if r["status"] == "OK")
    total_erros = sum(1 for r in resultados if r["status"] != "OK")

    ws.cell(row=row_resumo, column=1, value="RESUMO").font = Font(
        name="Arial", size=11, bold=True)
    resumo = [
        ("Petições geradas com sucesso:", total_ok),
        ("Registros com erro/aviso:", total_erros),
        ("Total de registros:", len(resultados)),
    ]
    for offset, (label, val) in enumerate(resumo, start=1):
        ws.cell(row=row_resumo + offset, column=1, value=label).font = Font(name="Arial", size=10)
        ws.cell(row=row_resumo + offset, column=2, value=val).font = Font(
            name="Arial", size=10, bold=True)

    # Congelar painel e filtro
    ws.freeze_panes = "A2"
    ultima_col = get_column_letter(len(colunas_header))
    ws.auto_filter.ref = f"A1:{ultima_col}{len(resultados) + 1}"

    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(caminho))
    log.info(f"  Relatório de conferência: {caminho.name}")
